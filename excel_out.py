from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
# ─────────────────────────────────────────────────────────────────
# STAGE 5: EXCEL OUTPUT
# ─────────────────────────────────────────────────────────────────
HEADER_BLUE   = "1F4E79"
HEADER_GREEN  = "1A5276"
ALT_BLUE      = "EBF3FB"
ALT_GREEN     = "E8F6F3"
FLAG_YELLOW   = "FFF2CC"
FLAG_RED      = "FDECEA"

def make_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def header_cell(cell, color):
    # Plain workbook output requested: white background and black text only.
    cell.font      = Font(bold=True, color="000000", name="Arial", size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = make_border()

def data_cell(cell, alt=False, flag=None, alt_color=ALT_BLUE):
    # Keep manual-curation and unmatched values as data only; do not colour cells.
    cell.font      = Font(name="Arial", size=10, color="000000")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = make_border()
    cell.fill      = PatternFill(fill_type=None)

def write_sheet(ws, headers, col_widths, rows_data, header_color, alt_color, flag_col=None):
    for col, h in enumerate(headers, 1):
        header_cell(ws.cell(row=1, column=col, value=h), header_color)
        ws.column_dimensions[get_column_letter(col)].width = col_widths[col-1]
    ws.row_dimensions[1].height = 40

    for row_i, row_data in enumerate(rows_data, 2):
        for col, val in enumerate(row_data, 1):
            data_cell(ws.cell(row=row_i, column=col, value=val))
        ws.row_dimensions[row_i].height = 30

    ws.freeze_panes = "A2"

def build_excel_multi(proc_rows: list[dict], finding_rows: list[dict], output_path: str, indication_rows: Optional[list[dict]] = None):
    """Build Excel with multiple procedure rows and their linked findings."""
    indication_rows = indication_rows or []
    wb = Workbook()

    # ── Sheet 1: Procedure ─────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Procedure"
    proc_headers = [
        "registration_id", "procedure_id", "procedure_finding_id",
        "procedure_date", "procedure_type", "procedure_location",
        "procedure_provider",
        "Procedure_completeness", "Reason_incomplete",
        "Intubation_distance", "Bowel_prep_given", "Bowel_prep_quality",
        "BBPS_right", "BBPS_transverse", "BBPS_left", "Total_BBPS_score",
        "Withdrawal_time", "Manual_curation", "Procedure_room",
        "Instrument_name", "Medication", "Tolerated_procedure",
        "Difficulty", "Difficulty_reason", "Complications",
        "Procedure_finding_2", "Procedure_finding_3", "Procedure_finding_4", "Procedure_finding_5",
        "Urgent_flag", "Extraction_discrepancy_flag",
        "pathology_report_text", "endoscopy_report_text"
    ]
    proc_widths = [18,16,24,16,20,22,24,20,18,20,22,20,12,16,12,16,16,24,16,18,16,18,32,20,20,20,20,20,20,16,26,100,100]

    rows_data = [[
        p.get("registration_id", ""),
        p["procedure_id"],
        p.get("procedure_finding_id", ""),
        p.get("procedure_date", ""),
        p.get("procedure_type", ""),
        p.get("procedure_location", ""),
        p.get("procedure_provider", ""),
        p["procedure_completeness"],
        p["reason_incomplete"],
        p["intubation_distance"],
        p["bowel_prep_given"],
        p["bowel_prep_quality"],
        p["bbps_right"],
        p["bbps_transverse"],
        p["bbps_left"],
        p["total_bbps_score"],
        p["withdrawal_time"],
        p["manual_curation"],
        p["procedure_room"],
        p["instrument_name"],
        p["medication"],
        p["tolerated_procedure"],
        p["difficulty"],
        p["difficulty_reason"],
        p["complications"],
        p.get("procedure_finding_2", p.get("other_findings", "None")),
        p.get("procedure_finding_3", ""),
        p.get("procedure_finding_4", ""),
        p.get("procedure_finding_5", ""),
        p.get("urgent_flag", "No"),
        p.get("extraction_discrepancy_flag", "No"),
        p.get("pathology_report_text", ""),
        p.get("endoscopy_report_text", ""),
    ] for p in proc_rows]

    write_sheet(ws1, proc_headers, proc_widths, rows_data, HEADER_BLUE, ALT_BLUE)

    # ── Sheet 2: Findings ──────────────────────────────────────────
    ws2 = wb.create_sheet("Findings")
    find_headers = [
        "finding_id", "registration_id", "procedure_id", "procedure_finding_id",
        "procedure_date", "procedure_type", "procedure_location",
        "Site", "Finding_morphology",
        "Number_categorical", "Size_mm",
        "Finding_intervention_performed", "Finding_intervention_completeness",
        "Type", "Diagnosis_group", "Dysplasia",
        "Path_site_matched", "Path_match_score", "Mixed_pathology",
        "Count_mismatch", "Negated", "Hedged", "Confidence",
        "Manual_curation", "Curation_reason",
        "Urgent_flag", "Extraction_discrepancy_flag",
        "pathology_report_text", "endoscopy_report_text"
    ]
    find_widths = [14,18,16,24,16,20,22,22,20,14,12,24,26,10,28,24,22,14,14,14,10,10,12,14,40,16,26,100,100]

    manual_cur_col = find_headers.index("Manual_curation") + 1

    frows_data = []
    for fr in finding_rows:
        frows_data.append([
            fr["finding_id"],
            fr.get("registration_id", ""),
            fr["procedure_id"],
            fr.get("procedure_finding_id", ""),
            fr.get("procedure_date", ""),
            fr.get("procedure_type", ""),
            fr.get("procedure_location", ""),
            fr["site"],
            fr["finding_morphology"],
            fr["number_categorical"],
            fr["size_mm"],
            fr["finding_intervention_performed"],
            fr["finding_intervention_completeness"],
            fr["type"],
            fr["diagnosis_group"],
            fr["dysplasia"],
            fr["path_site_matched"],
            fr["path_match_score"],
            fr["mixed_pathology"],
            fr["count_mismatch"],
            fr["negated"],
            fr["hedged"],
            fr["confidence"],
            fr["manual_curation"],
            fr["curation_reason"],
            fr.get("urgent_flag", "No"),
            fr.get("extraction_discrepancy_flag", "No"),
            fr.get("pathology_report_text", ""),
            fr.get("endoscopy_report_text", ""),
        ])

    write_sheet(ws2, find_headers, find_widths, frows_data,
                HEADER_GREEN, ALT_GREEN, flag_col=manual_cur_col)

    # ── Sheet 3: Indications ───────────────────────────────────────
    ws3 = wb.create_sheet("Indications")
    indication_headers = [
        "indication_id", "registration_id", "procedure_id", "procedure_finding_id",
        "procedure_date", "procedure_type", "procedure_location", "indication", "indication_nbcsp",
        "endoscopy_report_text"
    ]
    indication_widths = [14, 18, 16, 24, 16, 20, 22, 100, 100, 100]
    indication_data = [[
        row.get("indication_id", ""),
        row.get("registration_id", ""),
        row.get("procedure_id", ""),
        row.get("procedure_finding_id", ""),
        row.get("procedure_date", ""),
        row.get("procedure_type", ""),
        row.get("procedure_location", ""),
        row.get("indication", ""),
        row.get("indication_nbcsp", ""),
        row.get("endoscopy_report_text", ""),
    ] for row in indication_rows]
    write_sheet(ws3, indication_headers, indication_widths, indication_data, HEADER_BLUE, ALT_BLUE)

    # ── Legend ─────────────────────────────────────────────────────
    lr = len(finding_rows) + 4
    ws2.cell(row=lr, column=1, value="Legend:").font = Font(bold=True, name="Arial", size=10)

    wb.save(output_path)